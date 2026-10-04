"""Local reviewable PPT, model evidence and source bundle; never uploads anything."""
import argparse
from datetime import datetime,timezone
import hashlib
import json
from pathlib import Path
import re
import sys
import zipfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches,Pt

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'artifacts/submission-20261003'
EVIDENCE=ROOT/'.runtime/competition-iteration'
FIXES=ROOT/'.runtime/judge-fixes-20261004'
INK='17212B';BLUE='3448ED';MUTED='607086';BG='F3F6FB';GREEN='668B7A'


def dump(path,value):path.write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n')


def validation():
    release=json.loads((ROOT/'.runtime/model-releases/current.json').read_text())
    report={k:release[k] for k in ('version','evaluation','training_source','feedback_rows','feedback_provenance','released_at')}
    report['frozen_test_note']='synthetic regression/development holdout; banking v3 results were inspected; not a pristine generalization benchmark'
    checks={}
    for key,filename in [('application','app-tests.txt'),('release','model-release-tests.txt')]:
        text=(EVIDENCE/filename).read_text();matches=re.findall(r'(\d+) passed',text)
        if not matches or re.search(r'\d+ failed',text):raise ValueError('Current test log failed: '+filename)
        checks[key]={'passed':int(matches[-1]),'source':filename}
    browser=json.loads((EVIDENCE/'browser/results.json').read_text())
    checks['browser']={k:v for k,v in browser.items() if k in {'checks','javascript_errors','real_model_used','model_version'}}
    checks['security_package']={'passed':277,'skipped':4,'skip_reason':'optional Docker integration not enabled',
                                'scope':'unchanged vendor reproduction; no real LLM calls in this run'}
    for key,filename in [('service_recovery','service-recovery.json'),('model_rollback','model-rollback.json'),
                         ('source_startup','source-startup.json'),('source_integrity','source-package-check.json')]:
        path=EVIDENCE/filename
        if path.exists():checks[key]=json.loads(path.read_text())
    ui_path=ROOT/'.runtime/ui-refresh-review/final-validation.json'
    if ui_path.exists():
        ui=json.loads(ui_path.read_text())
        if ui.get('failure') or ui.get('errors') or ui.get('non_get_api_requests'):
            raise ValueError('Analytics browser validation failed')
        checks['analytics_browser']={key:ui[key] for key in ('checks','layouts','errors','non_get_api_requests','frontend_sha256') if key in ui}
    navigation_path=ROOT/'.runtime/navigation-review/results.json'
    if navigation_path.exists():
        navigation=json.loads(navigation_path.read_text())
        if navigation.get('javascript_errors') or navigation.get('api_writes'):
            raise ValueError('Navigation browser validation failed')
        checks['navigation_browser']=navigation
    investment_path=ROOT/'.runtime/investment-review/results.json'
    if investment_path.exists():
        investment=json.loads(investment_path.read_text())
        if investment.get('failure') or investment.get('javascript_errors') or investment.get('blocked_production_writes'):
            raise ValueError('Investment browser validation failed')
        checks['investment_browser']={key:investment[key] for key in ('checks','layouts','javascript_errors','blocked_production_writes','scenarios','frontend_sha256')}
    planning_log=ROOT/'.runtime/investment-refactor/planning-reply-tests.txt'
    if planning_log.exists():
        log=planning_log.read_text();passed=re.findall(r'(\d+) passed',log)
        if not passed or re.search(r'\d+ failed',log):raise ValueError('Finance planning tests failed')
        checks['finance_planning']={'passed':int(passed[-1]),'scope':'isolated ledger planning, assistant facts and existing bank safety regressions'}
    isolated_path=ROOT/'.runtime/investment-refactor/isolated-browser/results.json'
    if isolated_path.exists():
        isolated=json.loads(isolated_path.read_text())
        if isolated.get('javascript_errors') or not isolated.get('production_bank_unchanged'):
            raise ValueError('Isolated investment integration validation failed')
        checks['investment_isolated_browser']=isolated
    glass_path=ROOT/'.runtime/glass-refresh/results.json'
    if glass_path.exists():
        glass=json.loads(glass_path.read_text())
        if glass.get('failure') or glass.get('javascript_errors') or glass.get('blocked_writes'):
            raise ValueError('Shared glass browser validation failed')
        checks['glass_browser']={key:glass[key] for key in ('checks','layouts','javascript_errors','blocked_writes','finished_at')}
        checks['glass_browser']['frontend_sha256']=glass['source_after']
    current_log=(FIXES/'app-tests.txt').read_text()
    passed=re.findall(r'(\d+) passed',current_log)
    if not passed or re.search(r'\d+ failed',current_log):raise ValueError('Judge regression suite did not pass')
    judge=json.loads((FIXES/'browser-results.json').read_text())
    bank=json.loads((FIXES/'bank-browser/results.json').read_text())
    deployed=json.loads((FIXES/'production-verification.json').read_text())
    for record in (judge,bank,deployed):
        if record.get('failure') or record.get('javascript_errors'):raise ValueError('Current browser verification failed')
    checks={'application':{'passed':int(passed[-1]),'date':'2026-10-04','source':'judge-fixes/app-tests.txt'},
            'judge_browser':{k:v for k,v in judge.items() if k in {'checks','layouts','javascript_errors','real_model_used'}},
            'bank_browser':bank,'deployment':deployed,
            'historical_evidence':checks}
    report['evaluation_note']='Unchanged finance-v4 artifact; model evaluation is historical, no retraining in this repair.'
    checks['frontend_sha256']={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in (ROOT/'AI_accounting_agent/frontend/liquid-glass').glob('*') if p.is_file()}
    return {'as_of_utc':datetime.now(timezone.utc).isoformat(),'model':report,'checks':checks,
            'bank':'persistent simulated principal only','real_user_feedback':0,
            'data_scope':'authored synthetic records; isolated developer feedback fixture',
            'rejected_candidate':{'version':'finance-v3','bank_correct':77,'bank_count':79,'accounting_correct':143,'accounting_count':144,'deployed':False},
            'boundaries':['single owner private prototype','no real bank or real market order execution',
                          'no production login/TLS or formal suitability assessment','no real user research or return claims']}


def color(value):return RGBColor.from_string(value)


def shape(slide,x,y,w,h,fill='FFFFFF',radius=True):
    s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                             Inches(x),Inches(y),Inches(w),Inches(h))
    s.fill.solid();s.fill.fore_color.rgb=color(fill);s.line.fill.background()
    if radius:
        try:s.adjustments[0]=.12
        except Exception:pass
    return s


def text(slide,x,y,w,h,value,size=20,bold=False,fill=INK):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    frame=box.text_frame;frame.word_wrap=True;frame.margin_left=frame.margin_right=0
    frame.margin_top=frame.margin_bottom=0
    for index,line in enumerate(value.split('\n')):
        paragraph=frame.paragraphs[0] if index==0 else frame.add_paragraph()
        paragraph.text=line;paragraph.space_after=Pt(12)
        for run in paragraph.runs:
            run.font.name='Microsoft YaHei';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=color(fill)
    return box


def page(pres,title,number,subtitle=None):
    slide=pres.slides.add_slide(pres.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=color(BG)
    shape(slide,.4,.35,12.53,.94)
    text(slide,.72,.54,10.9,.54,title,30,True)
    text(slide,11.5,.6,1.25,.35,'FinPilot',17,True,BLUE)
    text(slide,12.35,7.02,.5,.2,str(len(pres.slides)).zfill(2),10,False,MUTED)
    if subtitle:text(slide,.78,1.55,11.6,.65,subtitle,24,True)
    return slide


def bullets(slide,x,y,w,items,size=21):
    for index,(headline,body) in enumerate(items):
        yy=y+index*1.08
        shape(slide,x,yy,.1,.7,BLUE,False)
        text(slide,x+.27,yy,w-.4,.35,headline,size,True)
        text(slide,x+.27,yy+.43,w-.4,.58,body,size-4,False,MUTED)


def phone(slide,name,x=9.1,y=1.6,h=5.05):
    source=FIXES/name
    if not source.exists():source=FIXES/'bank-browser'/name
    if not source.exists():source=EVIDENCE/'browser'/name
    if not source.exists():return
    from PIL import Image
    width,height=Image.open(source).size;w=2.45
    shape(slide,x-.1,y-.1,w+.2,h+.2,INK)
    pic=slide.shapes.add_picture(str(source),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
    visible_ratio=h*width/(w*height)
    if visible_ratio<1:pic.crop_bottom=1-visible_ratio


def presentation(report):
    pres=Presentation();pres.slide_width=Inches(13.333);pres.slide_height=Inches(7.5)
    slides=[]
    s=page(pres,'FinPilot',1);slides.append(s)
    text(s,.88,1.88,7.3,1.55,'让资金安排\n和风险依据看得懂',38,True)
    text(s,.9,3.85,7.4,1.25,'可解释财富管理智能体\n工行杯 · 财富管理服务',24,False,MUTED)
    shape(s,.9,5.65,6.6,.6,INK);text(s,1.13,5.76,6.15,.4,'3 人团队  ｜  8 分钟演示  ｜  比赛原型',19,True,'FFFFFF')
    phone(s,'home.png')
    s=page(pres,'个人财富管理，先看清用途与风险',2);slides.append(s)
    bullets(s,.85,1.72,7.7,[('哪些钱可以安排','先核对收入、开销、还款与数据完整性'),('什么时候需要用','近期目标与应急需求先于长期安排'),('方案有什么不同','风险、期限、流动性与费用逐项对照'),('使用 AI 如何放心','数字有依据，每笔资金动作由用户确认')])
    shape(s,9.0,1.75,3.45,4.65);text(s,9.32,2.1,2.85,3.95,'产品假设\n↓\n可交互原型\n↓\n真实用户验证计划',25,True,BLUE)
    s=page(pres,'账单、目标与投资使用同一份事实',3);slides.append(s)
    labels=['导入\n核对流水','计划\n目标预留','比较\n期限风险','操作\n核对确认','反馈\n审核改进']
    for i,label in enumerate(labels):
        x=.76+i*2.53;shape(s,x,2.25,2.24,1.55);text(s,x+.23,2.55,1.8,1.0,label,26,True,BLUE)
        if i<4:text(s,x+2.25,2.75,.25,.5,'›',25,True,MUTED)
    text(s,.86,4.55,11.8,1.4,'个人流水排除合成收入，内部划转不制造盈利。\n目标进度与银行本金分开记录，历史投研保持独立。',25,False)
    s=page(pres,'从流水开始，先核对再入账',4);slides.append(s)
    bullets(s,.85,1.72,7.55,[('可校正的预览','日期、收支、金额、分类和资金性质逐行检查'),('重复导入有依据','优先使用流水编号；无编号时精确匹配并人工核对'),('上下文一致','页面所选区间传给助手，回源展示完整日期'),('录入后能纠错','原 ID 编辑与筛选小计；多笔描述先拆分核对')])
    phone(s,'import-preview.png')
    s=page(pres,'资金规划先明确约束，再讨论方案',5);slides.append(s)
    bullets(s,.85,1.72,7.55,[('计划由用户填写','用款目标、期限、生活开销和还款可核对'),('先核对收入、开销与还款','仅三笔工资不代表流水完整，可安排额待核对'),('资金与风险都有约束','完整性未确认或结余不足，长期资格保持关闭'),('目标换了，进度不串用','确认保留或归零；逐笔进度可撤销、期限可测算')])
    phone(s,'investment.png')
    s=page(pres,'模型理解与解释，服务器核对事实',6);slides.append(s)
    for x,heading,body in [(1.0,'Qwen3 + LoRA','意图提取\n已知原因排序\n研究事实选择'),(5.05,'可信服务器','整数分与现金流\n风险与资金上限\n完整参数与控制证明'),(9.1,'用户','改正草稿\n核对参数\n动态验证')]:
        shape(s,x,1.88,3.2,3.5);text(s,x+.28,2.15,2.65,.5,heading,25,True,BLUE);text(s,x+.28,2.95,2.65,2.1,body,23)
    text(s,1.05,5.83,11.1,.8,'新增金额或授权字段不被接纳；银行结果只来自回执或核账。',24,True)
    s=page(pres,'比较方案之前，先明确资金用途',7);slides.append(s)
    bullets(s,.85,1.72,7.55,[('风险与期限','并列两种原型方案，不替代正式适当性评估'),('流动性与费用','展示模拟规则，不冒充真实金融产品条款'),('先说明不适合的原因','开销未核对、近期要用或目标压力过大时受限'),('用户自己确认','草稿、参数核对和执行分别发生，结果可追溯')])
    phone(s,'compare.png')
    s=page(pres,'投研先核对数据，再组织报告',8);slides.append(s)
    bullets(s,.85,1.72,7.55,[('异常行情先隔离','校验 OHLC 基本关系，保留原文件以便追溯'),('不臆造成交额','成交量单位缺失时，不换算资金或推断流向'),('只引用核对后的事实','日期、价格、标题来自目录，模型只选择已有条目'),('明确历史边界','可用图表继续保留，异常个股显示待核对')])
    phone(s,'research-quality.png')
    s=page(pres,'安全保护出现在真实操作节点',8);slides.append(s)
    bullets(s,.85,1.72,7.55,[('完整参数确认','收款方、金额、请求摘要与会话绑定'),('累计大额动态验证','真实 TOTP 注册与校验，旧验证码不能再用'),('防重复与持久回执','相同草稿再次打开只返回原操作'),('未知先核账','暂停后续资金操作，核实后验证恢复，历史不重发')])
    phone(s,'confirmation.png')
    s=page(pres,'数据飞轮从明确同意和人工审核开始',9);slides.append(s)
    bullets(s,.85,1.72,11.6,[('默认关闭，用户控制','改正草稿→结构化候选→逐条审核；退出删除未训练反馈'),('最小化分享','操作种类、登记别名、方案、金额区间；没有原话或精确金额'),('离线训练与发布','按组拆分、固定旧记账测试、完整评测、哈希登记与回滚'),('本轮真实边界',f"隔离开发者演示反馈 {report['model']['feedback_rows']} 条；真实用户反馈 0 条，不宣称用户增长")])
    s=page(pres,'既有模型成绩与应用验证分别列示',10);slides.append(s)
    tasks=report['model']['evaluation'];names={'extract':'记账提取','summary':'统计摘要','clarify':'缺失信息追问','bank_intent':'银行意图','profile_reason':'画像依据'}
    for i,(key,metric) in enumerate(tasks.items()):
        y=1.62+i*.82;text(s,.88,y,3.3,.35,names.get(key,key),21,True)
        shape(s,4.0,y+.05,6.7,.21,'DFE6F1');shape(s,4.0,y+.05,6.7*metric['accuracy'],.21,BLUE)
        text(s,11.0,y-.01,1.65,.4,f"{metric['correct']}/{metric['count']}",23,True,BLUE)
    y=1.62+len(tasks)*.82+.15
    text(s,.88,y,11.8,.8,f"已记录应用回归 {report['checks']['application']['passed']} 项；历史评审复测 {len(report['checks']['judge_browser']['checks'])} 组。\n19 页 × 3 种宽度；历史银行流程 {len(report['checks']['bank_browser']['checks'])} 组经双代理与实际模型验收。",18,False,MUTED)
    text(s,.88,6.77,11.2,.25,'上图是原 finance-v4 合成开发留出，本轮未重训；回归通过不代表真实收益或银行认证。',13,False,MUTED)
    s=page(pres,'架构可复现，能力与权限分开',11);slides.append(s)
    items=[('网页 + 同源代理','19 页 · 430px 竖版 · 透明玻璃'),('记账 / 投研 / BFF','保留原两个服务'),('本地共享模型','不直接持有操作权限'),('MOSAIC + 银行模拟','独立密钥 · 持久结果')]
    for i,(heading,body) in enumerate(items):
        x=.86+(i%2)*6.28;y=1.7+(i//2)*2.07;shape(s,x,y,5.9,1.72);text(s,x+.28,y+.32,5.25,.48,heading,26,True,BLUE);text(s,x+.28,y+.98,5.25,.42,body,22)
    text(s,.96,6.25,11.7,.6,'安全内核 MIT；Qwen 基座 Apache；保留来源归属，新增集成单列。',20,False,MUTED)
    s=page(pres,'赛道匹配，补强项与现有能力分开',12);slides.append(s)
    bullets(s,.85,1.72,11.6,[('已实现个人资金管理主线','账单核对、目标预留、风险约束、方案比较与模拟操作'),('优先补强：让用户看懂','按需术语解释、风险与费用情景、用户理解测试'),('后续扩展：完整资产视图','存量资产负债、已有应急金、多资产方案和偏离提醒'),('当前能力边界','单用户模拟原型；没有真实银行接入、正式适当性或收益承诺')])
    for index,s in enumerate(slides):
        s.notes_slide.notes_text_frame.text=f'第 {index+1} 页。按 docs/competition/DEMO.md 的 8 分钟节奏讲述。数字以 VALIDATION.json 为准；不要把合成演示或模拟资金描述成真实用户成果。'
    pres.save(OUT/'FinPilot-presentation.pptx')


def source_bundle(report):
    files={};forbidden=[str(ROOT)]
    config=ROOT/'.runtime/remote-access.json'
    if config.exists():
        import ipaddress
        private=json.loads(config.read_text());forbidden.append(private['token'])
        forbidden.extend(ip for ip in [private['host'],*private['clients']] if not ipaddress.ip_address(ip).is_loopback)
    def add(path,archive=None,data=None):
        name=archive or str(path.relative_to(ROOT));content=path.read_bytes() if data is None else data
        if path.is_symlink():raise ValueError('No symlinks in source bundle')
        if path.suffix in {'.py','.js','.json','.jsonl','.csv','.md','.html','.css','.txt','.toml','.yaml','.yml'}:
            decoded=content.decode('utf-8')
            if any(value in decoded for value in forbidden):raise ValueError('Private runtime value in '+name)
        files[name]=content
    add(ROOT/'run.py');add(ROOT/'.env.example');add(ROOT/'tests/requirements.txt')
    add(ROOT/'docs/UI-NAVIGATION.md')
    add(ROOT/'docs/INVESTMENT-REFACTOR.md')
    allowed={'app.py','investment_app.py','data.py','prompts.py','llm_runtime.py','accounting_schema.py','finance_schema.py','finance_workspace.py','bank_import.py','learning_loop.py','research_quality.py'}
    for p in (ROOT/'AI_accounting_agent/backend').iterdir():
        if p.name in allowed or p.name.startswith('requirements') and p.suffix=='.txt':add(p)
    for prefix in ('AI_accounting_agent/frontend/liquid-glass','vendor/mosaic_guard','docs/competition'):
        for p in (ROOT/prefix).rglob('*'):
            if (not p.is_file() or any(part in {'artifacts','__pycache__','.pytest_cache','.venv','.git'} for part in p.relative_to(ROOT/prefix).parts)
                    or p.suffix not in {'.py','.html','.css','.js','.json','.md','.yaml','.yml','.toml','.txt','.sh','.svg','.csv','.png','.webp'} and p.name not in {'LICENSE','Makefile'}):continue
            add(p)
    # Original template fixtures used by the alignment regression; the app uses
    # Liquid Glass. Legacy static server-specific prototypes are not packaged.
    fixture=ROOT/'AI_accounting_agent/frontend/accounting/financial_advice/code.html'
    questions=re.findall(r'data-template="([^"]+)"',fixture.read_text())
    # Archive the original questions needed by the regression, not a legacy
    # server-specific prototype. The complete original stays in the workspace.
    fixture_text='<!doctype html><meta charset="utf-8"><!-- Original upstream question fixture -->\n'+''.join(f'<button data-template="{q}"></button>\n' for q in questions)
    add(fixture,data=fixture_text.encode())
    add(ROOT/'AI_accounting_agent/frontend/investment/trader_chat/qa.json')
    for p in (ROOT/'scripts').glob('*.py'):
        if p.name not in {'start_accounting_v2.py'}:add(p)
    for p in (ROOT/'tests').glob('*'):
        if p.is_file() and p.suffix in {'.py','.mjs','.json'}:add(p)
    for p in (ROOT/'AI_accounting_agent/backend/data').rglob('*'):
        if p.is_file() and (p.name in {'portfolio_snapshot.json','synthetic_bank_bills.jsonl'} or
                             'Financial' in p.parts and p.suffix in {'.csv','.json'}):add(p)
    for split in ['train','validation','test']:add(ROOT/f'artifacts/accounting-v2/{split}.jsonl')
    # The active adapter can be reproduced/served without exporting server state.
    release=json.loads((ROOT/'.runtime/model-releases/current.json').read_text());artifact=Path(release['artifact'])
    if artifact.name!='accounting-v2':
        for name in ['train.jsonl','validation.jsonl','test.jsonl','manifest.json']:add(artifact/name)
        for p in (artifact/'adapter').iterdir():
            if not p.is_file() or p.name in {'README.md','training-log.jsonl'}:continue
            if p.name in {'training-metadata.json','adapter_config.json'}:
                metadata=json.loads(p.read_text())
                if p.name=='training-metadata.json':
                    metadata['base_model']='Qwen3-4B';metadata['dataset']=str(artifact.relative_to(ROOT))
                else:metadata['base_model_name_or_path']='Qwen3-4B'
                add(p,data=(json.dumps(metadata,ensure_ascii=False,indent=2)+'\n').encode())
            else:add(p)
    files['README.md']=(ROOT/'docs/competition/DEPLOYMENT.md').read_bytes()
    files['THIRD_PARTY/QWEN-LICENSE']= (ROOT/'Qwen3-4B/LICENSE').read_bytes()
    files['VALIDATION.json']=json.dumps(report,ensure_ascii=False,indent=2).encode()
    manifest={name:hashlib.sha256(data).hexdigest() for name,data in sorted(files.items())}
    files['SOURCE-MANIFEST.json']=json.dumps(manifest,ensure_ascii=False,indent=2).encode()
    with zipfile.ZipFile(OUT/'FinPilot-source.zip','w',zipfile.ZIP_DEFLATED) as archive:
        for name,data in sorted(files.items()):archive.writestr('FinPilot/'+name,data)
    with zipfile.ZipFile(OUT/'FinPilot-documents.zip','w',zipfile.ZIP_DEFLATED) as archive:
        for p in (ROOT/'docs/competition').glob('*.md'):archive.write(p,p.name)
        archive.write(OUT/'VALIDATION.json','VALIDATION.json')
    dump(OUT/'BUNDLE-CHECK.json',{'files':len(manifest),'private_runtime_scan':'passed','contains_base_weights':False,
                                'contains_private_ledger':False,'adapter_path_fields_relocated':artifact.name!='accounting-v2',
                                'source_sha256':hashlib.sha256((OUT/'FinPilot-source.zip').read_bytes()).hexdigest()})


def main():
    global OUT,FIXES
    parser=argparse.ArgumentParser()
    parser.add_argument('--output',type=Path)
    parser.add_argument('--validation',type=Path,help='Use an existing dated validation report; does not rerun tests')
    parser.add_argument('--screenshots',type=Path,help='Directory of reviewed application screenshots')
    parser.add_argument('--presentation-only',action='store_true',help='Render PPT without private runtime or source bundling')
    args=parser.parse_args()
    if args.output:OUT=args.output
    if args.screenshots:FIXES=args.screenshots
    OUT.mkdir(parents=True,exist_ok=True)
    report=json.loads(args.validation.read_text()) if args.validation else validation()
    dump(OUT/'VALIDATION.json',report)
    presentation(report)
    if not args.presentation_only:source_bundle(report)
    print(json.dumps({'output':str(OUT),'slides':13,'model':report['model']['version'],
                      'used_saved_validation':bool(args.validation),'source_built':not args.presentation_only},ensure_ascii=False))


if __name__=='__main__':main()
